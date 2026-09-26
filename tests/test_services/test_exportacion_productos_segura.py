"""V1.8-C: la exportación de productos (CSV y XLSX) no puede ejecutar fórmulas al abrirse y el ciclo
exportar -> reimportar devuelve exactamente los datos originales."""

import csv
import io
import itertools
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from db.repositorios import usuarios as repositorio_usuarios
from domain.celdas_seguras import desneutralizar_texto, neutralizar_texto
from domain.usuario import Usuario
from services import servicio_auth, servicio_exportacion, servicio_importacion, servicio_stock

NOMBRES_PELIGROSOS = ["=SUM(1+1)", "+cmd|'/C calc'!A0", "-2+3", "@SUM(A1)"]
NOMBRES_CON_APOSTROFO = ["'Cinco", "'=x", "''doble", "'-menos", "a'b"]
NOMBRES_NORMALES = ["Alfajor triple", "Ñandú 2=1", "x+y", "Té - verde"]


@pytest.fixture
def owner(base_datos_temporal):
    return repositorio_usuarios.crear_usuario(
        Usuario(
            nombre_usuario="duenio",
            nombre_completo="Dueño",
            password_hash=servicio_auth.hashear_password("clave-correcta-123", iteraciones=1000),
            rol="OWNER",
        )
    )


def _crear(nombre, codigo, costo=150, venta=300, stock=7):
    return servicio_stock.registrar_producto(codigo, nombre, costo, venta, stock_actual=stock)


def _celdas_csv(contenido: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(contenido.decode("utf-8-sig"), newline="")))


def _con_codigos_nuevos(contenido: bytes, sufijo: str) -> bytes:
    """El mismo CSV exportado, pero con otros códigos de barras (para importarlo como productos nuevos)."""
    filas = _celdas_csv(contenido)
    for fila in filas[1:]:
        fila[0] = fila[0] + sufijo
    buffer = io.StringIO()
    csv.writer(buffer).writerows(filas)
    return buffer.getvalue().encode("utf-8-sig")


class TestReglaInyectiva:
    ALFABETO = ["'", "=", "+", "-", "@", "\t", "\r", "a", " ", "1"]

    def test_desneutralizar_es_la_inversa_exacta_de_neutralizar_para_toda_cadena_corta(self):
        for largo in range(0, 5):
            for combinacion in itertools.product(self.ALFABETO, repeat=largo):
                texto = "".join(combinacion)
                assert desneutralizar_texto(neutralizar_texto(texto)) == texto, repr(texto)

    def test_la_neutralizacion_es_inyectiva(self):
        vistos: dict[str, str] = {}
        for largo in range(0, 5):
            for combinacion in itertools.product(self.ALFABETO, repeat=largo):
                texto = "".join(combinacion)
                salida = neutralizar_texto(texto)
                assert vistos.setdefault(salida, texto) == texto, (salida, texto, vistos[salida])

    def test_nada_que_empiece_como_formula_sale_sin_escapar(self):
        for combinacion in itertools.product(self.ALFABETO, repeat=3):
            salida = neutralizar_texto("".join(combinacion))
            assert not salida.startswith(("=", "+", "-", "@", "\t", "\r"))

    @pytest.mark.parametrize("texto", ["'Cinco", "a'b", "'", "'1", "' =x"])
    def test_un_apostrofo_inicial_que_no_es_nuestro_escape_se_respeta(self, texto):
        assert desneutralizar_texto(texto) == texto


class TestCsvDeProductos:
    def test_los_textos_peligrosos_salen_neutralizados_y_el_resto_igual(self, owner):
        todos = NOMBRES_PELIGROSOS + NOMBRES_NORMALES + NOMBRES_CON_APOSTROFO + ["\t=tab", "\r=cr"]
        for i, nombre in enumerate(todos):
            _crear(nombre, f"779000{i:04d}")
        _crear("Con codigo peligroso", "=1234")

        filas = _celdas_csv(servicio_exportacion.generar_csv_productos())

        por_codigo = {fila[0]: fila for fila in filas[1:]}
        for i, nombre in enumerate(todos):
            exportado = por_codigo[f"779000{i:04d}"][1]
            if nombre.startswith(("=", "+", "-", "@", "\t", "\r", "'")):
                assert exportado == "'" + nombre, nombre
            else:
                assert exportado == nombre, nombre
        assert "'=1234" in por_codigo  # el código de barras también es texto libre
        for fila in filas[1:]:
            for celda in fila:
                assert not celda.startswith(("=", "+", "-", "@", "\t", "\r")), celda

    def test_los_numeros_del_sistema_no_se_alteran(self, owner):
        _crear("Producto", "7790000000001", costo=12345, venta=99999, stock=42)

        fila = _celdas_csv(servicio_exportacion.generar_csv_productos())[1]

        assert fila[2:6] == ["123.45", "999.99", "42", "0"]

    def test_ida_y_vuelta_conserva_nombres_normales_peligrosos_y_con_apostrofo(self, owner):
        todos = NOMBRES_PELIGROSOS + NOMBRES_NORMALES + NOMBRES_CON_APOSTROFO
        for i, nombre in enumerate(todos):
            _crear(nombre, f"779000{i:04d}")
        exportado = servicio_exportacion.generar_csv_productos()

        resultado = servicio_importacion.procesar_archivo(
            "productos.csv", _con_codigos_nuevos(exportado, "N"), usuario_id=owner.id
        )

        assert (resultado.importados, resultado.errores) == (len(todos), [])
        for i, nombre in enumerate(todos):
            nuevo = servicio_stock.buscar_por_codigo_barras(f"779000{i:04d}N")
            assert nuevo.nombre == nombre, f"{nombre!r} volvió como {nuevo.nombre!r}"

    def test_reimportar_el_mismo_archivo_sobre_los_mismos_productos_no_cambia_nada(self, owner):
        todos = NOMBRES_PELIGROSOS + NOMBRES_NORMALES + NOMBRES_CON_APOSTROFO
        for i, nombre in enumerate(todos):
            _crear(nombre, f"779000{i:04d}")
        _crear("Con codigo peligroso", "=1234")
        exportado = servicio_exportacion.generar_csv_productos()

        resultado = servicio_importacion.procesar_archivo(
            "productos.csv", exportado, modo=servicio_importacion.MODO_CREAR_Y_ACTUALIZAR, usuario_id=owner.id
        )

        assert (resultado.actualizados, resultado.importados, resultado.errores) == (0, 0, [])
        assert resultado.sin_cambios == len(todos) + 1
        assert servicio_stock.buscar_por_codigo_barras("=1234").nombre == "Con codigo peligroso"
        assert sorted(p.nombre for p in servicio_stock.listar_todos() if p.codigo_barras.startswith("779000")) == sorted(todos)

    def test_un_apostrofo_legitimo_de_un_archivo_armado_a_mano_no_se_toca(self, owner):
        contenido = (
            "codigo_barras,nombre,precio_costo,precio_venta,stock_actual,stock_minimo,categoria,unidad_medida\n"
            "A1,'Cinco,1.00,2.00,0,0,,UNIDAD\n"
            "A2,''doble,1.00,2.00,0,0,,UNIDAD\n"
        ).encode("utf-8")

        servicio_importacion.procesar_archivo("x.csv", contenido, usuario_id=owner.id)

        # `''doble` cumple la forma de NUESTRO escape (apóstrofo + apóstrofo): vuelve como `'doble`.
        assert servicio_stock.buscar_por_codigo_barras("A1").nombre == "'Cinco"
        assert servicio_stock.buscar_por_codigo_barras("A2").nombre == "'doble"

    @pytest.mark.parametrize("nombre, esperado", [("\t=tab", "=tab"), ("\r=cr", "=cr")])
    def test_espacios_iniciales_tab_y_cr_salen_neutralizados_y_la_importacion_los_recorta_como_siempre(
        self, owner, nombre, esperado
    ):
        _crear(nombre, "7790009999")
        exportado = servicio_exportacion.generar_csv_productos()
        assert _celdas_csv(exportado)[1][1] == "'" + nombre

        servicio_importacion.procesar_archivo("x.csv", _con_codigos_nuevos(exportado, "N"), usuario_id=owner.id)

        # La importación siempre recortó los espacios de los extremos (comportamiento previo, sin cambios).
        assert servicio_stock.buscar_por_codigo_barras("7790009999N").nombre == esperado


class TestXlsxDeProductos:
    def test_el_riesgo_existe_openpyxl_guarda_como_formula_un_texto_con_igual(self):
        libro = Workbook()
        libro.active.append(["=1+1", "+1", "-1", "@x"])

        assert [c.data_type for c in libro.active[1]] == ["f", "s", "s", "s"]

    def test_todo_texto_se_exporta_como_texto_y_con_el_valor_exacto(self, owner):
        todos = NOMBRES_PELIGROSOS + NOMBRES_NORMALES + NOMBRES_CON_APOSTROFO + ["\t=tab", "\r=cr"]
        for i, nombre in enumerate(todos):
            _crear(nombre, f"779000{i:04d}")
        _crear("Con codigo peligroso", "=1234")

        contenido = servicio_exportacion.generar_xlsx_productos()

        hoja = load_workbook(io.BytesIO(contenido)).active
        celdas = [c for fila in hoja.iter_rows() for c in fila]
        assert {c.data_type for c in celdas if c.value is not None} == {"s"}
        nombres = {fila[0].value: fila[1].value for fila in hoja.iter_rows(min_row=2)}
        for i, nombre in enumerate(todos):
            # Sin cambiar el valor: nada que deshacer al reimportar. (El XML de XLSX normaliza un CR a salto de línea:
            # es propio del formato, no de la protección, y tampoco lo vuelve una fórmula.)
            assert nombres[f"779000{i:04d}"] == nombre.replace(chr(13), chr(10))
        assert "=1234" in nombres
        with zipfile.ZipFile(io.BytesIO(contenido)) as zip_xlsx:
            assert b"<f>" not in zip_xlsx.read("xl/worksheets/sheet1.xml")

    def test_ida_y_vuelta_del_xlsx_no_cambia_ningun_producto(self, owner):
        todos = NOMBRES_PELIGROSOS + NOMBRES_NORMALES + NOMBRES_CON_APOSTROFO
        for i, nombre in enumerate(todos):
            _crear(nombre, f"779000{i:04d}")
        contenido = servicio_exportacion.generar_xlsx_productos()

        resultado = servicio_importacion.procesar_archivo(
            "productos.xlsx", contenido, modo=servicio_importacion.MODO_CREAR_Y_ACTUALIZAR, usuario_id=owner.id
        )

        assert (resultado.actualizados, resultado.errores, resultado.sin_cambios) == (0, [], len(todos))
        assert sorted(p.nombre for p in servicio_stock.listar_todos()) == sorted(todos)
