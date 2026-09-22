"""Pruebas de integración de services.servicio_exportacion contra una base
de datos SQLite real (temporal y aislada, ver tests/conftest.py).

Cubre generación de CSV/XLSX, que solo se exporten productos activos, y
el ciclo completo exportar -> reimportar (ver
tests/test_services/test_servicio_importacion.py para las pruebas de
`servicio_importacion` en sí, que acá se reutiliza para probar el
round-trip, no se vuelve a probar su comportamiento interno)."""

import csv
import io

from openpyxl import load_workbook

from db.conexion import obtener_conexion
from services import servicio_categorias, servicio_exportacion, servicio_importacion, servicio_stock


def _filas_csv(contenido: bytes) -> list[list[str]]:
    texto = contenido.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(texto)))


def _filas_xlsx(contenido: bytes) -> list[tuple]:
    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    return list(libro.active.iter_rows(values_only=True))


class TestGenerarCsvProductos:
    def test_catalogo_vacio_solo_tiene_encabezado(self, base_datos_temporal):
        filas = _filas_csv(servicio_exportacion.generar_csv_productos())

        assert filas == [list(servicio_exportacion.COLUMNAS_EXPORTACION)]

    def test_encabezado_coincide_con_las_columnas_que_espera_la_importacion(self, base_datos_temporal):
        filas = _filas_csv(servicio_exportacion.generar_csv_productos())

        encabezado = filas[0]
        for columna in servicio_importacion.COLUMNAS_REQUERIDAS:
            assert columna in encabezado

    def test_contenido_de_un_producto(self, base_datos_temporal):
        categoria = servicio_categorias.crear_categoria("Golosinas")
        servicio_stock.registrar_producto(
            codigo_barras="7790000000001",
            nombre="Alfajor",
            precio_costo_centavos=15050,
            precio_venta_centavos=25000,
            stock_actual=10,
            stock_minimo=2,
            categoria_id=categoria.id,
            unidad_medida="UNIDAD",
        )

        filas = _filas_csv(servicio_exportacion.generar_csv_productos())

        assert filas[0] == list(servicio_exportacion.COLUMNAS_EXPORTACION)
        assert filas[1] == ["7790000000001", "Alfajor", "150.50", "250.00", "10", "2", "Golosinas", "UNIDAD"]

    def test_solo_incluye_productos_activos(self, base_datos_temporal):
        # Baja lógica directa (activo=0, sin borrar la fila): así la
        # prueba cubre el filtro `WHERE activo = 1` en sí, no el caso
        # trivial de un producto que ya ni siquiera existe en la tabla
        # (que es lo que haría `servicio_stock.eliminar_producto` para
        # un producto sin ventas -- ver `db.repositorios.productos`).
        servicio_stock.registrar_producto("7790000000001", "Activo", 100, 200)
        producto_baja = servicio_stock.registrar_producto("7790000000002", "De baja", 100, 200)
        with obtener_conexion() as conexion:
            conexion.execute("UPDATE productos SET activo = 0 WHERE id = ?", (producto_baja.id,))

        filas = _filas_csv(servicio_exportacion.generar_csv_productos())

        codigos = [fila[0] for fila in filas[1:]]
        assert codigos == ["7790000000001"]

    def test_producto_sin_categoria_exporta_columna_vacia(self, base_datos_temporal):
        servicio_stock.registrar_producto("7790000000001", "Sin categoría", 100, 200)

        filas = _filas_csv(servicio_exportacion.generar_csv_productos())

        assert filas[1][6] == ""  # columna "categoria"

    def test_conserva_nombre_de_categoria_dada_de_baja(self, base_datos_temporal):
        categoria = servicio_categorias.crear_categoria("Descontinuada")
        servicio_stock.registrar_producto(
            "7790000000001", "Producto viejo", 100, 200, categoria_id=categoria.id
        )
        servicio_categorias.eliminar_categoria(categoria.id)

        filas = _filas_csv(servicio_exportacion.generar_csv_productos())

        assert filas[1][6] == "Descontinuada"

    def test_codificacion_utf8_sig_conserva_acentos(self, base_datos_temporal):
        servicio_stock.registrar_producto("7790000000001", "Alfajor de dulce de leche - Edición Ñoño", 100, 200)

        contenido = servicio_exportacion.generar_csv_productos()

        assert contenido.startswith(b"\xef\xbb\xbf")  # BOM UTF-8
        filas = _filas_csv(contenido)
        assert filas[1][1] == "Alfajor de dulce de leche - Edición Ñoño"


class TestGenerarXlsxProductos:
    def test_catalogo_vacio_solo_tiene_encabezado(self, base_datos_temporal):
        filas = _filas_xlsx(servicio_exportacion.generar_xlsx_productos())

        assert filas == [tuple(servicio_exportacion.COLUMNAS_EXPORTACION)]

    def test_contenido_de_un_producto(self, base_datos_temporal):
        categoria = servicio_categorias.crear_categoria("Bebidas")
        servicio_stock.registrar_producto(
            codigo_barras="7790000000002",
            nombre="Gaseosa",
            precio_costo_centavos=20000,
            precio_venta_centavos=35000,
            stock_actual=5,
            stock_minimo=1,
            categoria_id=categoria.id,
            unidad_medida="LITRO",
        )

        filas = _filas_xlsx(servicio_exportacion.generar_xlsx_productos())

        assert filas[0] == tuple(servicio_exportacion.COLUMNAS_EXPORTACION)
        # Todas las columnas se escriben como texto (incluso stock, que no
        # es dinero): mismo camino que el CSV, y evita representar precios
        # como `float` en la celda (ver docstring de `_fila_exportable`).
        assert filas[1] == ("7790000000002", "Gaseosa", "200.00", "350.00", "5", "1", "Bebidas", "LITRO")

    def test_solo_incluye_productos_activos(self, base_datos_temporal):
        servicio_stock.registrar_producto("7790000000001", "Activo", 100, 200)
        producto_baja = servicio_stock.registrar_producto("7790000000002", "De baja", 100, 200)
        with obtener_conexion() as conexion:
            conexion.execute("UPDATE productos SET activo = 0 WHERE id = ?", (producto_baja.id,))

        filas = _filas_xlsx(servicio_exportacion.generar_xlsx_productos())

        assert len(filas) == 2  # encabezado + 1 producto


class TestCompatibilidadExportarReimportar:
    """Round-trip real: exportar -> borrar los productos originales ->
    reimportar el archivo exportado, sin editarlo -> confirmar que los
    datos vuelven a quedar iguales. Prueba explícitamente la garantía
    pedida (exportar -> descargar -> importar sin adaptación manual)."""

    def _productos_de_prueba(self):
        categoria = servicio_categorias.crear_categoria("Golosinas")
        servicio_stock.registrar_producto(
            codigo_barras="7790000000001",
            nombre="Alfajor",
            precio_costo_centavos=15050,
            precio_venta_centavos=25099,
            stock_actual=10,
            stock_minimo=2,
            categoria_id=categoria.id,
            unidad_medida="UNIDAD",
        )
        servicio_stock.registrar_producto(
            codigo_barras="7790000000002",
            nombre="Gaseosa 2L",
            precio_costo_centavos=20000,
            precio_venta_centavos=35000,
            stock_actual=5,
            stock_minimo=1,
            unidad_medida="LITRO",
        )

    def _borrar_todos_los_productos(self):
        with obtener_conexion() as conexion:
            conexion.execute("DELETE FROM productos")

    def test_csv_exportado_se_reimporta_sin_adaptacion(self, base_datos_temporal):
        self._productos_de_prueba()
        contenido = servicio_exportacion.generar_csv_productos()
        self._borrar_todos_los_productos()

        resultado = servicio_importacion.procesar_archivo("productos.csv", contenido)

        assert resultado.importados == 2
        assert resultado.duplicados == 0
        assert resultado.errores == []

        alfajor = servicio_stock.buscar_por_codigo_barras("7790000000001")
        assert alfajor.nombre == "Alfajor"
        assert alfajor.precio_costo_centavos == 15050
        assert alfajor.precio_venta_centavos == 25099
        assert alfajor.stock_actual == 10
        assert alfajor.stock_minimo == 2
        assert alfajor.unidad_medida == "UNIDAD"
        categoria_reimportada = servicio_categorias.obtener_por_id(alfajor.categoria_id)
        assert categoria_reimportada.nombre == "Golosinas"

        gaseosa = servicio_stock.buscar_por_codigo_barras("7790000000002")
        assert gaseosa.precio_venta_centavos == 35000
        assert gaseosa.unidad_medida == "LITRO"
        assert gaseosa.categoria_id is None

    def test_xlsx_exportado_se_reimporta_sin_adaptacion(self, base_datos_temporal):
        self._productos_de_prueba()
        contenido = servicio_exportacion.generar_xlsx_productos()
        self._borrar_todos_los_productos()

        resultado = servicio_importacion.procesar_archivo("productos.xlsx", contenido)

        assert resultado.importados == 2
        assert resultado.duplicados == 0
        assert resultado.errores == []

        alfajor = servicio_stock.buscar_por_codigo_barras("7790000000001")
        assert alfajor.precio_costo_centavos == 15050
        assert alfajor.precio_venta_centavos == 25099
        assert alfajor.stock_actual == 10
