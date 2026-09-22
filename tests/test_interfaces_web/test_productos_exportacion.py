"""Pruebas HTTP end-to-end de la exportación de productos
(`GET /productos/exportar/csv` y `/xlsx`).

Complementan `tests/test_services/test_servicio_exportacion.py` (que ya
prueba el contenido generado): acá se prueba la ruta en sí -- control
de acceso, headers de descarga, y el caso de catálogo vacío contra la
app real."""

from domain.usuario import Usuario
from interfaces.web.auth import NOMBRE_COOKIE_SESION
from services import servicio_auth, servicio_stock

from ._asgi_cliente import solicitud


def _cookies_para_rol(rol: str, nombre_usuario: str) -> dict[str, str]:
    from db.repositorios import usuarios as repositorio_usuarios

    repositorio_usuarios.crear_usuario(
        Usuario(
            nombre_usuario=nombre_usuario,
            nombre_completo="Usuario de prueba",
            password_hash=servicio_auth.hashear_password("clave-correcta-123", iteraciones=1000),
            rol=rol,
        )
    )
    token = servicio_auth.iniciar_sesion(nombre_usuario, "clave-correcta-123").token
    return {NOMBRE_COOKIE_SESION: token}


def _cookies_owner() -> dict[str, str]:
    return _cookies_para_rol("OWNER", "ana")


def _cookies_cashier() -> dict[str, str]:
    return _cookies_para_rol("CASHIER", "carla")


class TestExportarCsv:
    def test_owner_puede_descargar(self, base_datos_temporal):
        servicio_stock.registrar_producto("7790000000001", "Alfajor", 100, 200)
        cookies = _cookies_owner()

        respuesta = solicitud("GET", "/productos/exportar/csv", cookies=cookies)

        assert respuesta.status == 200
        assert (respuesta.header("content-type") or "").startswith("text/csv")
        content_disposition = respuesta.header("content-disposition") or ""
        assert 'attachment; filename="productos_' in content_disposition
        assert content_disposition.endswith('.csv"')

    def test_cashier_no_puede_descargar(self, base_datos_temporal):
        cookies = _cookies_cashier()

        respuesta = solicitud("GET", "/productos/exportar/csv", cookies=cookies)

        assert respuesta.status == 403

    def test_sin_sesion_no_puede_descargar(self, base_datos_temporal):
        respuesta = solicitud("GET", "/productos/exportar/csv")

        assert respuesta.status in (401, 303)

    def test_catalogo_vacio_devuelve_solo_encabezado(self, base_datos_temporal):
        cookies = _cookies_owner()

        respuesta = solicitud("GET", "/productos/exportar/csv", cookies=cookies)

        assert respuesta.status == 200
        lineas = respuesta.texto.strip("﻿").strip().splitlines()
        assert len(lineas) == 1
        assert "codigo_barras" in lineas[0]


class TestExportarXlsx:
    def test_owner_puede_descargar(self, base_datos_temporal):
        servicio_stock.registrar_producto("7790000000001", "Alfajor", 100, 200)
        cookies = _cookies_owner()

        respuesta = solicitud("GET", "/productos/exportar/xlsx", cookies=cookies)

        assert respuesta.status == 200
        assert respuesta.header("content-type") == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        content_disposition = respuesta.header("content-disposition") or ""
        assert 'attachment; filename="productos_' in content_disposition
        assert content_disposition.endswith('.xlsx"')

    def test_cashier_no_puede_descargar(self, base_datos_temporal):
        cookies = _cookies_cashier()

        respuesta = solicitud("GET", "/productos/exportar/xlsx", cookies=cookies)

        assert respuesta.status == 403

    def test_contenido_es_un_xlsx_valido(self, base_datos_temporal):
        import io

        from openpyxl import load_workbook

        servicio_stock.registrar_producto("7790000000001", "Alfajor", 100, 200)
        cookies = _cookies_owner()

        respuesta = solicitud("GET", "/productos/exportar/xlsx", cookies=cookies)

        libro = load_workbook(io.BytesIO(respuesta.cuerpo), read_only=True)
        filas = list(libro.active.iter_rows(values_only=True))
        assert filas[0][0] == "codigo_barras"
        assert filas[1][0] == "7790000000001"
